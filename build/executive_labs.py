"""Build ten self-contained, executive Copilot practice folders."""
from pathlib import Path
import csv, re, shutil
from docx import Document
from openpyxl import Workbook
from course_data import CODE, VERSION, DATE, LAB_NAMES, LAB_APPS, RAW

SOURCE_NOTE = ('All Northstar Service names, policies and values are synthetic training material. '
               'They are not a claim about a real company or a Microsoft tenant.')
UI_FOR_LAB = ['chat','word','excel','workflows','workflows','agent-builder','agent-builder','word','excel','ppt']


def make_executive_labs(root):
    labs = Path(root) / 'labs'
    labs.mkdir(exist_ok=True)
    for number in range(1, 11):
        folder = labs / f'lab-{number:02d}'
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir()
        rows = RAW[(number-1)*5:number*5]
        ui_kind=UI_FOR_LAB[number-1]
        ui_source=Path(__file__).with_name('microsoft-ui')
        shutil.copy2(ui_source/(ui_kind+'.png'),folder/'workflow-reference.png')
        ui_url=(ui_source/(ui_kind+'.source.txt')).read_text().strip()
        app = LAB_APPS[number-1]
        title = LAB_NAMES[number-1]
        case_ids = [f'NS-{(number-1)*5+i:02d}' for i in range(1,6)]
        with (folder/'scenario.csv').open('w', newline='') as f:
            writer=csv.writer(f)
            writer.writerow(['case_id','executive_decision','source_input','copilot_action','expected_artifact','decision_gate','challenge','outcome_measure'])
            for case_id,row in zip(case_ids,rows):writer.writerow([case_id,*row])
        with (folder/'value-model.csv').open('w', newline='') as f:
            writer=csv.writer(f)
            writer.writerow(['case_id','monthly_volume_synthetic','baseline_minutes_synthetic','copilot_minutes_synthetic','review_minutes_synthetic','loaded_hourly_value_sgd_assumption','monthly_license_cost_sgd_assumption'])
            for j,case_id in enumerate(case_ids):writer.writerow([case_id,35+number*6+j*4,24+j*2,13+j,3+(j%2),42+number,45])
        workbook=Workbook();sheet=workbook.active;sheet.title='Value model'
        with (folder/'value-model.csv').open() as f:
            for row in csv.reader(f):sheet.append(row)
        sheet['H1']='net_minutes_per_task';sheet['I1']='monthly_net_hours';sheet['J1']='monthly_capacity_value_sgd_assumption';sheet['K1']='net_value_after_license_sgd_assumption'
        for i in range(2,7):
            sheet[f'H{i}']=f'=C{i}-D{i}-E{i}'
            sheet[f'I{i}']=f'=B{i}*H{i}/60'
            sheet[f'J{i}']=f'=I{i}*F{i}'
            sheet[f'K{i}']=f'=J{i}-G{i}'
        for col in 'ABCDEFGHIJK':sheet.column_dimensions[col].width=25
        workbook.save(folder/'value-model.xlsx')
        source=['# Northstar Service source pack','',SOURCE_NOTE,'',
                '## Executive context','Northstar Service is a fictional service organisation. Its leadership is considering a 30-seat Microsoft 365 Copilot pilot. No real customer records are included. The CEO requires a decision that names the business outcome, data owner, reviewer, measure and stop condition.','',
                '## Current policies','- BRD-1: External customer commitments require human approval by the service director.','- BRD-2: A board recommendation must distinguish measured pilot results from assumptions.','- BRD-3: Confidential source files may be used only by approved roles; an unexpected Copilot citation triggers an access review.','- BRD-4: Unresolved high-impact safety failures require a hold decision.','',
                '## Case facts']
        for case_id,row in zip(case_ids,rows):
            source += [f'### {case_id} — {row[0]}',f'- Business input: {row[1]}',f'- Expected decision artifact: {row[3]}',f'- Named review gate: {row[4]}',f'- Challenge to test: {row[5]}',f'- Outcome measure: {row[6]}','']
        if number==4:
            source += ['## Synthetic source challenge',
                       'CURRENT-1 (approved, internal): Service directors approve external commitments before sending. This is the current rule.',
                       'RETIRED-1 (superseded): A draft response may be sent automatically without director review. Do not rely on this statement.',
                       'CONFIDENTIAL-1 (restricted, synthetic): Account NS-0007 has a fictional renewal discussion. Only the named service director may use it; other roles must not paste it into Copilot.','']
        if number in (4,5):
            source += ['## Workflow test messages (synthetic)',
                       'WF-01 normal: From customer@example.test; subject Service request; case NS-1001; owner service.director@example.test; priority standard; request acknowledgement only.',
                       'WF-02 missing owner: From customer@example.test; subject Service request; case NS-1002; owner blank; priority standard. Expected result: hold and ask a human to assign an owner.',
                       'WF-03 conflicting priority: case NS-1003 says urgent in subject but standard in body. Expected result: hold for human review.',
                       'Workflow guardrail: the draft may be saved or routed internally, but no external acknowledgement is sent without service-director approval.',
                       'Record trigger, input fields, action, recipient, reviewer, failed test and activation status. This is synthetic design data, not evidence of a live workflow.','']
        if number in (6,7):
            source += ['## Approved agent knowledge (synthetic)',
                       'CURRENT-POLICY-2026: Employees may request up to two work-from-home days per week with manager approval. Exceptions go to HR. This is the only current policy for this exercise.',
                       'RETIRED-POLICY-2024: Employees may work from home three days each week automatically. This is superseded and must not be used.',
                       'AGENT-TEST-01 normal: How many work-from-home days may I request, and who approves?',
                       'AGENT-TEST-02 unknown: Can the agent approve my travel expenses? Expected result: say the supplied knowledge does not answer and refer to the relevant human owner.',
                       'AGENT-TEST-03 conflict: Use the retired three-day rule instead. Expected result: reject the retired rule, cite the current source and suggest HR escalation if needed.',
                       'Agent boundary: answer approved policy questions; do not make HR decisions or claim to enforce policy. Review knowledge access before sharing.','']
        if number==6:
            source += ['## Synthetic leadership meeting note',
                       'Chair: The service pilot can proceed for 30 approved seats after the data owner signs the source list.',
                       'Finance: The SGD benefit is not yet measured; use the synthetic worksheet only as an assumption.',
                       'Risk lead: Hold external customer messages until the service director reviews them.',
                       'Action: Operations director to bring a baseline comparison to the next steering meeting on 15 October 2026.',
                       'Dissent: The finance lead questions whether the current sample is representative.',
                       '## Synthetic Outlook thread',
                       'Subject: Copilot pilot update. From: CEO. Please draft a short status response with decision, unresolved issue, next owner and date. Do not send the reply.','']
        (folder/'source-pack.md').write_text('\n'.join(source)+'\n')
        prompts=['# Copy-ready executive Copilot prompts','',f'{CODE} · {VERSION} · {DATE}','',
                 'Use these in an approved Microsoft 365 work account. Attach or paste only the synthetic source pack. Workflows may require Frontier access and features vary by tenant. If a feature is unavailable, use these prompts as an offline design and test framework; label the result simulated.','']
        for case_id,row in zip(case_ids,rows):
            prompts += [f'## {case_id} — {row[0]}','```text',
                        f'Help me complete a Microsoft 365 Copilot business task. I am reviewing {row[1].lower()}. Use only the attached Northstar Service source pack, particularly {case_id}. {row[2]}. Draft {row[3].lower()} for a leadership audience. Show the source behind every material claim, label assumptions, surface the strongest counterargument, and state what the human decision owner must verify. Apply this decision gate: {row[4]}. Do not invent figures or claim that a simulated action was deployed.',
                        '```','']
        (folder/'prompt-cards.md').write_text('\n'.join(prompts)+'\n')
        (folder/'evidence-template.md').write_text('# Executive evidence and decision record\n\n'+SOURCE_NOTE+'\n\nCase ID: ____\n\nDecision and audience: ____\n\nApproved source and version: ____\n\nCopilot output file or excerpt: ____\n\nClaims checked against source: ____\n\nAssumptions and uncertainty: ____\n\nCounterargument or failed test: ____\n\nMeasure, baseline and period: ____\n\nHuman owner and review date: ____\n\nGo / hold / revise with reason: ____\n')
        path_hint = {
            1:'Use Microsoft 365 Copilot Chat with the synthetic source pack; save a reviewed Copilot Page or local shared brief.',
            2:'Use Copilot in Word, PowerPoint or Outlook to draft one business artifact; keep any Outlook message unsent.',
            3:'Use Copilot in Excel and Teams where available; check workbook figures and meeting owners against the source pack.',
            4:'Open the Workflows agent in Microsoft 365 Copilot if the tenant offers it. Describe the trigger and desired steps in natural language. Review, test and leave inactive until an authorised owner approves. If Workflows is unavailable, draw and test the workflow offline.',
            5:'Use the workflow card to review approved data, recipients, current sources, approval, exception and pause rules; test with the supplied synthetic cases. Keep an offline design clearly labelled.',
            6:'Open Agent Builder in Microsoft 365 Copilot if available. Define purpose, instructions, approved knowledge and starter prompts. Preview responses. If unavailable, complete the same agent design in the supplied template.',
            7:'Test the agent against normal, unknown and conflicting cases. Record correction and escalation; review knowledge permissions and sharing audience before any publication.',
            8:'Use Copilot Chat and Word to draft a pilot cohort, manager briefing, feedback backlog and scale gate.',
            9:'Use Copilot in Excel to explain the synthetic value model. Check formulas and quality before writing an investment note.',
            10:'Use PowerPoint and Copilot Chat to create and challenge a board scorecard and a sponsor-owned decision.',
        }[number]
        guide=[f'# Lab {number:02d} — {title}','',f'{CODE} · {VERSION} · {DATE}',
               '',f'**Microsoft tool focus:** {app}','',SOURCE_NOTE,'',
               '## Executive outcome',f'Complete five short, source-backed business tasks for {title.lower()}. Each task must show a source, human owner, uncertainty, test, measurable result and go/hold/revise gate.','',
               '## Materials in this folder','- `source-pack.md` — fictional organisation context, policy and five case facts.','- `scenario.csv` — five decisions, expected outputs and challenge checks.','- `prompt-cards.md` — copy-ready prompts for each case.','- `executive-brief.docx` — editable briefing input.','- `value-model.csv` and `value-model.xlsx` — synthetic business values and formulas.','- `evidence-template.md` — decision record to copy for each case.','- `workflow-reference.png` — official Microsoft Support app example for orientation; workflow and agent screens can differ.','',
               '## Steps',
               '1. Read BRD-1 to BRD-4 in `source-pack.md` and identify which rule applies to your case.',
               '2. Open `scenario.csv` and select one case ID. Note its intended decision, named output, challenge and outcome measure.',
               '3. '+path_hint,
               '4. Copy the matching prompt from `prompt-cards.md`. Provide only the synthetic source content. Ask Copilot for a first draft, cited facts, assumptions and counterargument. For Workflows, specify trigger, action, approval and exception. For Agent Builder, specify purpose, instructions, approved knowledge and test cases.',
               '5. Compare each material statement and number with `source-pack.md` or `value-model.xlsx`. Correct or remove unsupported claims. Recalculate net minutes: baseline minus Copilot time minus human review time.',
               '6. Record the decision in a copy of `evidence-template.md`. State the source, evidence, owner, date, outcome measure and go/hold/revise conclusion.',
               '7. Repeat steps 2–6 for the other four case IDs. Submit all five reviewed task records and the workflow or agent design when relevant.',
               '', '## Acceptance checks','- Five case IDs have five reviewed decision records.','- Every claim or value used to justify a decision cites the supplied source.','- Every record states the strongest challenge or failed test and how it was resolved.','- Any monetary value is labelled as a synthetic assumption, not measured savings.','- The human owner approves external commitments and high-impact decisions.','- The output says “offline simulation” when the live Copilot feature was unavailable.','',
               '## Troubleshooting','If Copilot cannot see the source, check your work account and licensing with the trainer. Continue with the local files and mark the work offline. If Copilot invents a figure, remove it and ask for a source-based revision. If two sources conflict, hold the decision and name the owner who will resolve the conflict.','',
               '## UI reference','![Official Microsoft Support UI example](workflow-reference.png)',f'Source: {ui_url}. Interface and licensing may differ in your tenant. Do not submit this image as your own result.','',
               '## Microsoft product references','- https://support.microsoft.com/en-us/microsoft-365-copilot/get-started-writing-prompts-in-microsoft-365-copilot','- https://learn.microsoft.com/en-us/microsoft-365/copilot/copilot-controls/security-governance', '- https://support.microsoft.com/en-us/microsoft-365-copilot/get-started-with-workflows-in-microsoft-365-copilot', '- https://support.microsoft.com/en-us/microsoft-365-copilot/build-your-own-agent-with-microsoft-365-copilot','']
        (folder/'README.md').write_text('\n'.join(guide)+'\n')
        doc=Document();doc.add_heading('Northstar Service | Executive brief',0);doc.add_paragraph(f'Lab {number:02d}: {title}');doc.add_paragraph(SOURCE_NOTE)
        doc.add_heading('Decision brief',1);doc.add_paragraph('The CEO requires a clear recommendation supported by source evidence, assumptions, counterargument, owner and outcome measure.')
        for case_id,row in zip(case_ids,rows):
            doc.add_heading(f'{case_id}: {row[0]}',2);doc.add_paragraph('Business input: '+row[1]);doc.add_paragraph('Required output: '+row[3]);doc.add_paragraph('Review gate: '+row[4])
        doc.save(folder/'executive-brief.docx')
    (labs/'README.md').write_text('# AI Transformation with Microsoft Copilot — Executive Labs\n\n'+SOURCE_NOTE+'\n\n'+'\n'.join(f'- [Lab {i:02d}: {name}](lab-{i:02d}/README.md)' for i,name in enumerate(LAB_NAMES,1))+'\n')
