"""Build the trainer-facing Facilitator Guide and Assessment Plan."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from course_data import CODE, DATE, OUTCOMES, TITLE, TOPICS, VERSION


ROOT = Path(__file__).resolve().parents[1] / "courseware"
NAVY = RGBColor(19, 47, 83)


def new_doc(kind):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 22), ("Heading 1", 13), ("Heading 2", 10)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
    doc.add_picture(str(Path(__file__).with_name("provider-logo.png")), width=Inches(1.65))
    doc.add_heading(kind, 0)
    doc.add_paragraph(TITLE, style="Subtitle")
    doc.add_paragraph(f"{CODE}  •  {VERSION}  •  {DATE}  •  WSQ  •  3 days / 24 hours")
    doc.add_paragraph("Tertiary Infotech Academy Pte Ltd  •  UEN 201200696W")
    doc.add_paragraph("Trainer controlled document. Use synthetic learner data only.")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = f"Tertiary Infotech Academy  •  {CODE}  •  {kind}"
    return doc


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    return t


def build_fg():
    doc = new_doc("Facilitator Guide")
    doc.add_heading("Executive delivery intent", 1)
    doc.add_paragraph("Guide senior leaders through a synthetic Northstar Service decision. The course uses Microsoft Copilot directly in Chat, Word, PowerPoint, Outlook, Teams and Excel, then tests the human judgment required to turn drafts into responsible business decisions. AB-730 informs direct business use; AB-100 informs strategy and outcome judgment. This WSQ course is not Microsoft certification preparation.")
    doc.add_heading("Prepare before class", 1)
    bullets(doc, [
        "Check the current v2.0 slide deck, Learner Guide, Lesson Plan, Labs 01–10, and candidate WA/PP papers against TGS-2024044051.",
        "Open each lab's source pack, prompts, executive brief and value-model workbook. Confirm the synthetic figures and formulas are readable.",
        "If licensed Microsoft 365 Copilot is available, use an approved work account. Keep the offline evidence path ready; never portray a simulation as a live tenant result.",
        "Ask executives to bring a business decision, success measure and human owner. Use only synthetic case data in shared exercises."
    ])
    doc.add_heading("Three-day facilitation sequence", 1)
    table(doc, ["Block", "Business work and Copilot use", "Trainer checkpoint"], [
        ("Day 1 · LO1", "Strategy mandate, Copilot Chat decision brief, Word memo, PowerPoint board story; Labs 01–03", "Challenge each claim, baseline and board recommendation."),
        ("Day 2 · LO2", "Source boundaries, responsible use, governance charter; Labs 04–05", "Escalate an unsupported or confidential claim; name approval owner."),
        ("Day 2 · LO3 bridge", "Teams and Outlook executive cadence; Lab 06", "Verify decision, dissent, recipient and external commitment."),
        ("Day 3 · LO3–LO4", "Integration options, adoption plan, Excel value model, board scorecard; Labs 07–10", "Compare as-is, extend, build; label observed values and assumptions."),
        ("Day 3 · assessment", "WA 60 minutes and PP 60 minutes", "Use only candidate papers; keep marking guides controlled.")
    ])
    doc.add_page_break()
    doc.add_heading("Learning outcomes and evidence", 1)
    for i, outcome in enumerate(OUTCOMES, 1):
        doc.add_paragraph(f"LO{i}: {outcome}", style="List Number")
    doc.add_heading("Facilitation moves", 1)
    bullets(doc, [
        "Demonstrate a Copilot draft, then ask leaders which statements they would sign, question or remove. Show the source and human edit.",
        "For each lab, require a decision artifact with business outcome, source, assumption, counterargument, owner and go/hold/revise gate.",
        "Treat usage as an adoption indicator, not an outcome. Check comparable work, review time, quality and risk before estimating value.",
        "Use the five synthetic cases in each lab. An unavailable feature becomes a clearly labelled offline decision exercise, not a fabricated screenshot."
    ])
    doc.add_heading("Assessment handoff", 1)
    doc.add_paragraph("The Written Assessment covers K1–K3. The Practical Performance has four tasks covering A1–A5: executive transformation mandate, governance charter, integration decision and value-led board recommendation. Apply the current controlled marking guide and record criterion-level C/NYC evidence.")
    doc.save(ROOT / "FG-AI-Transformation-with-Microsoft-Copilot.docx")


def build_ap():
    doc = new_doc("Assessment Plan")
    doc.add_heading("Assessment overview", 1)
    doc.add_paragraph("Assess the four approved WSQ outcomes with one Written Assessment (WA) and one Practical Performance (PP). The final two hours are part of the 24-hour course. Candidate work uses synthetic Northstar Service data and may use an approved Copilot work account or a clearly labelled offline equivalent.")
    table(doc, ["Method", "Duration", "Candidate evidence", "Decision"], [
        ("WA", "60 min", "Three applied executive questions covering K1–K3", "C/NYC against controlled marking guide"),
        ("PP", "60 min", "Four decision artifacts covering A1–A5", "C/NYC against controlled evidence rubric")
    ])
    doc.add_heading("Outcome and criterion map", 1)
    table(doc, ["Outcome", "Knowledge evidence", "Performance evidence"], [
        ("LO1 · Implementation and process", "K1 · outcome, pilot, Copilot use and baseline", "A1–A2 · mandate and implementation plan"),
        ("LO2 · Security controls", "K2 · sources, verification and human approval", "A3 · responsible-use and control charter"),
        ("LO3 · Integration alignment", "K3 · use-as-is, extend or build trade-offs", "A4 · integration decision and workflow test"),
        ("LO4 · Optimization", "K3 · adoption, quality, risk and value", "A5 · value model and board go/hold recommendation")
    ])
    doc.add_heading("Administration", 1)
    bullets(doc, [
        "Confirm candidate identity and issue only the current v2.0 WA/PP candidate papers. Keep answer keys and assessor decisions outside public learner locations.",
        "State time, open-book rules, permitted synthetic data, tool access and submission format before starting. An offline path must be labelled simulated.",
        "Provide approved reasonable adjustments without changing the standard. Record the adjustment and the evidence still required.",
        "Collect each candidate's own source-backed decision artifacts; do not coach toward a particular recommendation."
    ])
    doc.add_page_break()
    doc.add_heading("Evidence and decision", 1)
    bullets(doc, [
        "Check the source behind each material claim and figure. Reject unsupported savings or a simulated action described as deployed.",
        "Check decision owner, approval gate, counterargument, risk response and a measurable outcome for every relevant task.",
        "Award Competent only when every required criterion is met. Otherwise record Not Yet Competent with the specific gap and reassessment action.",
        "Store submissions, assessor decisions and reassessment records in the approved controlled location under provider retention rules."
    ])
    doc.add_heading("Related documents", 1)
    doc.add_paragraph("Use the v2.0 Trainer Slides, Learner Guide, Lesson Plan and ten executive labs for instruction. The WA and PP candidate papers are the only learner-facing assessment instruments; marking guides are assessor-only.")
    doc.save(ROOT / "AP-AI-Transformation-with-Microsoft-Copilot.docx")

if __name__ == "__main__":
    build_fg()
    build_ap()
